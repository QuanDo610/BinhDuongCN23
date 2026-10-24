import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Độc lập - Tự do - Hạnh phúc")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.bold = True

p_line = doc.add_paragraph()
p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_line.paragraph_format.space_before = Pt(0)
p_line.paragraph_format.space_after = Pt(2)
# Draw horizontal line using VML or Drawing or exact em-dashes
# Let's test w:drawing with a shape line
line_xml = parse_xml(r'''
<w:r %s>
    <w:drawing>
        <wp:inline distT="0" distB="0" distL="0" distR="0">
            <wp:extent cx="1500000" cy="12000"/>
            <wp:docPr id="1" name="Line 1"/>
            <wp:cNvGraphicFramePr/>
            <a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/main">
                    <a:xfrm>
                        <a:off x="0" y="0"/>
                        <a:ext cx="1500000" cy="12000"/>
                    </a:xfrm>
                    <a:prstGeom prst="rect">
                        <a:avLst/>
                    </a:prstGeom>
                    <a:solidFill>
                        <a:srgbClr val="000000"/>
                    </a:solidFill>
                </a:graphicData>
            </a:graphic>
        </wp:inline>
    </w:drawing>
</w:r>
''' % nsdecls('w', 'wp'))
p_line._p.append(line_xml)
doc.save("scratch/test_drawing.docx")
print("Saved test_drawing.docx")
