import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Độc lập - Tự do - Hạnh phúc")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.bold = True

# Option 1: Drawing line exact length
p_line1 = doc.add_paragraph()
p_line1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_line1.paragraph_format.space_before = Pt(0)
p_line1.paragraph_format.space_after = Pt(4)
# In EMUs: 1 inch = 914400 EMUs. Text width is approx 1.75 - 1.85 inches = 1600000 EMUs. Height = 10000 EMUs (approx 0.8pt)
def add_horizontal_line(paragraph, width_inches=1.8, height_pt=0.8):
    cx = int(width_inches * 914400)
    cy = int(height_pt * 12700)
    line_xml = parse_xml(r'''
    <w:r %s>
        <w:drawing>
            <wp:inline distT="0" distB="0" distL="0" distR="0">
                <wp:extent cx="%d" cy="%d"/>
                <wp:docPr id="1" name="Line"/>
                <wp:cNvGraphicFramePr/>
                <a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                    <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/main">
                        <a:xfrm>
                            <a:off x="0" y="0"/>
                            <a:ext cx="%d" cy="%d"/>
                        </a:xfrm>
                        <a:prstGeom prst="rect">
                            <a:avLst/>
                        </a:prstGeom>
                        <a:solidFill>
                            <a:srgbClr val="000000"/>
                        </a:solidFill>
                    </a:graphicData>
                </a:graphic>
            </wp:inline>
        </w:drawing>
    </w:r>
    ''' % (nsdecls('w', 'wp'), cx, cy, cx, cy))
    paragraph._p.append(line_xml)

add_horizontal_line(p_line1, width_inches=1.85, height_pt=0.75)

# Option 2: Em dashes continuous string
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Độc lập - Tự do - Hạnh phúc")
r2.font.name = "Times New Roman"
r2.font.size = Pt(13)
r2.bold = True

p_line2 = doc.add_paragraph()
p_line2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_line2.paragraph_format.space_before = Pt(0)
p_line2.paragraph_format.space_after = Pt(4)
r_dash = p_line2.add_run("—" * 22) # 22 em-dashes
r_dash.font.name = "Times New Roman"
r_dash.font.size = Pt(8.5)
r_dash.bold = True

doc.save("scratch/compare.docx")
print("Saved compare.docx")
