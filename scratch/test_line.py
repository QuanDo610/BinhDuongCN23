import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
section = doc.sections[0]
section.top_margin = Inches(0.79)
section.bottom_margin = Inches(0.79)
section.left_margin = Inches(1.18)
section.right_margin = Inches(0.59)

# Test paragraph border
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Độc lập - Tự do - Hạnh phúc")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.bold = True

# Add paragraph bottom border
pPr = p._p.get_or_add_pPr()
pBdr = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="6" w:space="2" w:color="000000"/></w:pBdr>' % nsdecls('w'))
pPr.append(pBdr)

doc.save("scratch/test_border.docx")
print("Saved test_border.docx")
