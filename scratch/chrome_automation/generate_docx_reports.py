import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho ô trong bảng"""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề trong ô bảng (dxd)"""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(15, 23, 42)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_header_block(doc):
    # Header bảng 2 cột: Cơ quan bên trái, Tiêu ngữ bên phải
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Cột trái
    cell_left = tbl.cell(0, 0)
    cell_left.width = Inches(3.2)
    p_cq = cell_left.paragraphs[0]
    p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_cq.add_run("VĂN PHÒNG ĐĂNG KÝ ĐẤT ĐAI TP.HCM\n")
    format_run(r1, size_pt=10, bold=True)
    r2 = p_cq.add_run("CHI NHÁNH SỐ 23\n")
    format_run(r2, size_pt=10.5, bold=True, color_rgb=(2, 132, 199))
    r3 = p_cq.add_run("Số: ...... /BC-CN23")
    format_run(r3, size_pt=9.5, italic=True)

    # Cột phải
    cell_right = tbl.cell(0, 1)
    cell_right.width = Inches(3.8)
    p_tn = cell_right.paragraphs[0]
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p_tn.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    format_run(r4, size_pt=10, bold=True)
    r5 = p_tn.add_run("Độc lập - Tự do - Hạnh phúc\n")
    format_run(r5, size_pt=10.5, bold=True)
    r6 = p_tn.add_run("Thủ Dầu Một, ngày 28 tháng 09 năm 2026")
    format_run(r6, size_pt=9.5, italic=True)

    # Xóa viền bảng header
    for cell in tbl.rows[0].cells:
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(6)

def add_signature_block(doc, author_name="Đỗ Trọng Minh Quân"):
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(14)
    
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell_left = tbl.cell(0, 0)
    cell_left.width = Inches(3.5)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_nl = p_left.add_run("Nơi nhận:\n")
    format_run(r_nl, size_pt=10, bold=True, italic=True)
    r_nl2 = p_left.add_run("- Như trên (báo cáo);\n- Chị Bùi Thị Minh Nguyệt (phối hợp);\n- Lưu: Hồ sơ kỹ thuật IT.")
    format_run(r_nl2, size_pt=9.5, italic=True)

    cell_right = tbl.cell(0, 1)
    cell_right.width = Inches(3.5)
    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_right.add_run("NGƯỜI LẬP BÁO CÁO\n\n\n\n\n")
    format_run(r_title, size_pt=11, bold=True)
    r_name = p_right.add_run(author_name)
    format_run(r_name, size_pt=11.5, bold=True, color_rgb=(15, 23, 42))

    for cell in tbl.rows[0].cells:
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)

# ==============================================================================
# TẠO FILE 1: BÁO CÁO KHẢO SÁT HỆ THỐNG OERP
# ==============================================================================
def create_report_1(output_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

    add_header_block(doc)

    # Tiêu đề báo cáo
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("BÁO CÁO KHẢO SÁT & ĐÁNH GIÁ TOÀN DIỆN\nHỆ THỐNG PHẦN MỀM OERP ĐIỀU HÀNH TÁC NGHIỆP ĐẤT ĐAI")
    format_run(r_t, size_pt=14, bold=True, color_rgb=(3, 105, 161))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("(Hệ thống phục vụ chuyên môn tại địa chỉ web: https://vpdkddtdm.vn)")
    format_run(r_sub, size_pt=10.5, italic=True)

    # Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.paragraph_format.space_after = Pt(12)
    p_kg.paragraph_format.left_indent = Inches(0.5)
    r_kg = p_kg.add_run("Kính gửi: ")
    format_run(r_kg, size_pt=11.5, bold=True)
    r_kg2 = p_kg.add_run("Chị Nguyễn Thị Bích Tuyền\n")
    format_run(r_kg2, size_pt=11.5, bold=True, color_rgb=(2, 132, 199))
    r_dg = p_kg.add_run("Đồng gửi: ")
    format_run(r_dg, size_pt=11, bold=True)
    r_dg2 = p_kg.add_run("Chị Bùi Thị Minh Nguyệt (Bộ phận chuyên môn)")
    format_run(r_dg2, size_pt=11, italic=True)

    # Lời dẫn
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.25
    p_intro.paragraph_format.space_after = Pt(8)
    r_in = p_intro.add_run(
        "Thực hiện công tác rà soát, đánh giá hiện trạng các phần mềm và hạ tầng công nghệ thông tin phục vụ tác nghiệp tại Chi nhánh, "
        "tôi đã tiến hành kiểm tra kỹ thuật chi tiết đối với hệ thống phần mềm OERP đang vận hành tại địa chỉ "
    )
    format_run(r_in, size_pt=11)
    r_in2 = p_intro.add_run("https://vpdkddtdm.vn")
    format_run(r_in2, size_pt=11, bold=True)
    r_in3 = p_intro.add_run(". Nay xin trân trọng báo cáo kết quả khảo sát kỹ thuật cụ thể như sau:")
    format_run(r_in3, size_pt=11)

    # 1. Mục tiêu và nghiệp vụ
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("I. MỤC TIÊU NGHIỆP VỤ VÀ VAI TRÒ CỦA HỆ THỐNG")
    format_run(r_h1, size_pt=12, bold=True, color_rgb=(3, 105, 161))

    p_m1 = doc.add_paragraph()
    p_m1.paragraph_format.line_spacing = 1.2
    p_m1.paragraph_format.space_after = Pt(4)
    r_m1 = p_m1.add_run(
        "Hệ thống OERP tại Chi nhánh là phần mềm tác nghiệp chuyên sâu, đóng vai trò số hóa toàn bộ chu trình xử lý hồ sơ địa chính từ khâu tiếp nhận của người dân cho đến khi trao trả kết quả. Cụ thể phục vụ 4 mục tiêu trọng tâm:"
    )
    format_run(r_m1, size_pt=11)

    bullets1 = [
        ("Số hóa quy trình Một Cửa liên thông:", " Tiếp nhận hồ sơ cấp sổ lần đầu, chuyển nhượng, tách thửa, đo vẽ trích lục địa chính của cả tổ chức và cá nhân."),
        ("Kiểm soát mốc thời gian & Chống trễ hạn:", " Ghi nhận chính xác thời điểm luân chuyển hồ sơ giữa các bộ phận (Một cửa -> Đo đạc -> Thẩm tra -> Thuế -> Trả kết quả) nhằm phục vụ chỉ số Cải cách hành chính (CCHC)."),
        ("Tra cứu ngăn chặn giao dịch tức thì:", " Hỗ trợ cán bộ kiểm tra ngay thửa đất có đang bị tranh chấp, kê biên thi hành án hay cấm chuyển nhượng trước khi lập biên nhận."),
        ("Tương tác tự động với người dân:", " Tích hợp cổng SMS Brandname tự động gửi tin nhắn báo tiến độ và mời người dân đến nhận kết quả đúng hẹn.")
    ]
    for b_title, b_desc in bullets1:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        format_run(r_bt, size_pt=11, bold=True)
        r_bd = p_b.add_run(b_desc)
        format_run(r_bd, size_pt=11)

    # 2. Khảo sát Sitemap & 9 Phân hệ
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("II. BẢN ĐỒ CHỨC NĂNG (SITEMAP) & CÁC PHÂN HỆ NGHIỆP VỤ")
    format_run(r_h2, size_pt=12, bold=True, color_rgb=(3, 105, 161))

    p_sm_desc = doc.add_paragraph()
    p_sm_desc.paragraph_format.space_after = Pt(6)
    r_smd = p_sm_desc.add_run("Qua khảo sát thực tế qua tài khoản nghiệp vụ, hệ thống gồm 9 phân hệ chức năng khép kín:")
    format_run(r_smd, size_pt=11)

    # Bảng 9 phân hệ
    tbl_ph = doc.add_table(rows=1, cols=3)
    tbl_ph.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ph.autofit = False

    headers = ["STT", "Phân hệ & Đường dẫn", "Chức năng & Tương tác nghiệp vụ chính"]
    col_widths = [Inches(0.6), Inches(2.3), Inches(4.1)]
    for i, h in enumerate(headers):
        cell = tbl_ph.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "E0F2FE")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        format_run(r, size_pt=10.5, bold=True, color_rgb=(3, 105, 161))

    subsystems = [
        ("1", "Trang Chủ\n/vptdm/trangchu", "Danh sách các hồ sơ được phân bổ cho cá nhân cán bộ đăng nhập cần thụ lý trong ngày."),
        ("2", "Tiếp Nhận Hồ Sơ\n/vptdm/tiepnhanhoso", "Nhập thông tin số tờ, số thửa, số nhà, diện tích. In biên nhận Một Cửa, xuất Word, cấu hình lịch đo đạc, tra cứu ngăn chặn giao dịch và gửi SMS tự động."),
        ("3", "Xử Lý Hồ Sơ\n/vptdm/xulyhoso", "Phân công lịch đo đạc địa chính, chức năng thuế, chuyển hồ sơ nội bộ giữa các khâu, giao phiếu nhanh, xuất Excel theo dõi."),
        ("4", "Sửa Hồ Sơ\n/vptdm/suahoso", "Điều chỉnh thông tin nhân thân, thửa đất khi có sai lệch; quản lý danh mục tài liệu yêu cầu người dân bổ sung."),
        ("5", "Trả Hồ Sơ\n/vptdm/trahoso", "Bàn giao sổ đỏ/kết quả trích đo cho khách hàng, ký nhận kết quả, chức năng xuất và nhập danh sách qua Excel."),
        ("6", "Tra Cứu Hồ Sơ\n/vptdm/tracuuhoso", "Tra cứu tiến độ hồ sơ theo mã biên nhận Một Cửa hoặc tên khách hàng/công ty. Xem danh sách ảnh scan và bản vẽ trích lục."),
        ("7", "Quản Lý Trạng Thái\n/vptdm/quanlytrangthaihoso", "Chức năng dành cho lãnh đạo/tổ trưởng: Hoàn thành cưỡng chế, Hủy hồ sơ sai thủ tục, Khôi phục hồ sơ cần xử lý lại."),
        ("8", "Dashboard Trích Đo\n/vptdm/dashboard-trichdo", "Biểu đồ trực quan giám sát tiến độ xử lý hồ sơ trích đo theo từng điểm quy trình đo đạc để phát hiện điểm nghẽn."),
        ("9", "Thống Kê Tổng Hợp\n/vptdm/thongketonghop", "Bộ lọc theo tham số báo cáo, thống kê tỷ lệ đúng hạn/trễ hạn, xuất dữ liệu chi tiết ra file Excel phục vụ báo cáo lãnh đạo.")
    ]

    for row_idx, data in enumerate(subsystems):
        row = tbl_ph.add_row()
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                format_run(r, size_pt=10, bold=True)
            elif col_idx == 1:
                lines = text.split("\n")
                r1 = p.add_run(lines[0] + "\n")
                format_run(r1, size_pt=10, bold=True, color_rgb=(15, 23, 42))
                if len(lines) > 1:
                    r2 = p.add_run(lines[1])
                    format_run(r2, size_pt=9, italic=True, color_rgb=(2, 132, 199))
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                format_run(r, size_pt=10)

    # Đặt viền cho bảng
    for row in tbl_ph.rows:
        for cell in row.cells:
            tcPr = cell._element.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
            tcPr.append(tcBorders)

    # 3. Kiến trúc kỹ thuật & Nhà cung cấp
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("III. KIẾN TRÚC KỸ THUẬT & THÔNG TIN NHÀ CUNG CẤP TÊN MIỀN/SERVER")
    format_run(r_h3, size_pt=12, bold=True, color_rgb=(3, 105, 161))

    bullets2 = [
        ("Frontend (Giao diện):", " Được xây dựng bằng công nghệ Single Page Application (Angular SPA) kết hợp Bootstrap 4.5.2 và jQuery. Giao diện chạy hoàn toàn trên trình duyệt cán bộ, tải rất nhanh."),
        ("Backend (Máy chủ xử lý API):", " Máy chủ Windows Server chạy Microsoft IIS với nền tảng ASP.NET (phiên bản v4.0.30319). Địa chỉ máy chủ API: apisql.vpdkddtdm.vn."),
        ("Cơ sở dữ liệu (Database):", " Sử dụng Microsoft SQL Server (MSSQL). Dữ liệu được truy vấn trực tiếp qua các thủ tục lưu sẵn (Stored Procedures: A1011SELEDN, HT2000, postA1012DangNhapA)."),
        ("Tên miền & Nhà cung cấp:", " Qua tra cứu chính thức tại Hệ thống Quản lý Tên miền Quốc gia (VNNIC), tên miền vpdkddtdm.vn đang do cá nhân ông TRẦN HỒ BẢO HOÀNG đứng tên đăng ký qua Công ty Cổ phần GMO-Z.com RUNSYSTEM (Đăng ký ngày 07/07/2026, hết hạn 07/07/2027)."),
        ("Chi phí thuê ngoài hiện tại:", " Cơ quan đang chi trả 800.000 đồng/tháng (tương đương 9.600.000 đồng/năm) để thuê chỗ đặt server và duy trì hệ thống bên ngoài.")
    ]
    for b_title, b_desc in bullets2:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        format_run(r_bt, size_pt=11, bold=True)
        r_bd = p_b.add_run(b_desc)
        format_run(r_bd, size_pt=11)

    add_signature_block(doc)
    doc.save(output_path)
    print(f"✅ Đã tạo xong Báo cáo 1: {output_path}")

# ==============================================================================
# TẠO FILE 2: BÁO CÁO PHƯƠNG ÁN DI DỜI SERVER NỘI BỘ (0 ĐỒNG)
# ==============================================================================
def create_report_2(output_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

    add_header_block(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("BÁO CÁO ĐỀ XUẤT PHƯƠNG ÁN DI DỜI HỆ THỐNG OERP\nVỀ PHÒNG MÁY CHỦ NỘI BỘ CHI NHÁNH (CHI PHÍ 0 ĐỒNG)")
    format_run(r_t, size_pt=14, bold=True, color_rgb=(3, 105, 161))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("(Thực hiện mục tiêu cắt giảm chi phí 800k/tháng, tự chủ CSDL đất đai và bảo đảm an toàn thông tin)")
    format_run(r_sub, size_pt=10.5, italic=True)

    # Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.paragraph_format.space_after = Pt(12)
    p_kg.paragraph_format.left_indent = Inches(0.5)
    r_kg = p_kg.add_run("Kính gửi: ")
    format_run(r_kg, size_pt=11.5, bold=True)
    r_kg2 = p_kg.add_run("Chị Nguyễn Thị Bích Tuyền\n")
    format_run(r_kg2, size_pt=11.5, bold=True, color_rgb=(2, 132, 199))
    r_dg = p_kg.add_run("Đồng gửi: ")
    format_run(r_dg, size_pt=11, bold=True)
    r_dg2 = p_kg.add_run("Chị Bùi Thị Minh Nguyệt (Bộ phận chuyên môn)")
    format_run(r_dg2, size_pt=11, italic=True)

    # Lời dẫn
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.25
    p_intro.paragraph_format.space_after = Pt(8)
    r_in = p_intro.add_run(
        "Theo định hướng của đơn vị về việc tăng cường bảo mật dữ liệu chuyên ngành đất đai, không để cơ sở dữ liệu trên máy chủ thuê ngoài "
        "và thực hiện tiết kiệm tối đa ngân sách thường xuyên (tiêu chí tận dụng hạ tầng sẵn có, cắt giảm chi phí thuê ngoài 800.000 đ/tháng), "
        "tôi xin đề xuất phương án kỹ thuật chi tiết để di dời toàn bộ hệ thống OERP về phòng máy chủ của Chi nhánh như sau:"
    )
    format_run(r_in, size_pt=11)

    # 1. Cơ sở thực tiễn
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("I. KHẢ NĂNG TỰ CHỦ HẠ TẦNG VỚI CHI PHÍ ĐẦU TƯ 0 ĐỒNG")
    format_run(r_h1, size_pt=12, bold=True, color_rgb=(3, 105, 161))

    p_cs = doc.add_paragraph()
    p_cs.paragraph_format.space_after = Pt(4)
    r_cs = p_cs.add_run("Hiện nay Chi nhánh số 23 đã được trang bị phòng máy chủ đạt chuẩn với các thiết bị vật lý chuyên dụng đang hoạt động:")
    format_run(r_cs, size_pt=11)

    bullets1 = [
        ("Máy chủ chuyên dụng lưu trữ HP StoreEasy 1460 (IP: 10.17.2.100):", " Đang chứa sẵn folder Dataweb và HSQ-TDM, dung lượng ổ đĩa lớn, rất thích hợp làm nơi lưu trữ các file scan sổ đỏ và bản vẽ trích đo."),
        ("Máy chủ hiệu năng cao HP ProLiant DL360 Gen9 (IP: 10.17.2.125):", " Cấu hình vi xử lý và RAM rất mạnh, dư thừa công suất để cài đặt máy chủ Web IIS và cơ sở dữ liệu Microsoft SQL Server cho hệ thống OERP."),
        ("Thiết bị lưu trữ mạng NAS Synology DS423+ (IP: 10.17.2.150):", " Có 4 khay ổ cứng, sẵn sàng thiết lập lịch sao lưu tự động (auto-backup) toàn bộ cơ sở dữ liệu OERP vào 23h00 hàng đêm."),
        ("Đường truyền mạng & An ninh:", " Có sẵn 02 đường truyền Internet cáp quang VNPT và tường lửa bảo mật phần cứng FortiGate 61F.")
    ]
    for b_title, b_desc in bullets1:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(3)
        r_bt = p_b.add_run(b_title)
        format_run(r_bt, size_pt=11, bold=True)
        r_bd = p_b.add_run(b_desc)
        format_run(r_bd, size_pt=11)

    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.space_before = Pt(4)
    p_concl.paragraph_format.space_after = Pt(6)
    r_c = p_concl.add_run("=> Đánh giá tài chính: Đơn vị KHÔNG PHẢI MUA SẮM THÊM BẤT KỲ MÁY CHỦ MỚI NÀO. Khi chuyển về cơ quan, đơn vị sẽ cắt giảm được ngay 9.600.000 đ/năm tiền thuê máy chủ bên ngoài.")
    format_run(r_c, size_pt=11, bold=True, color_rgb=(5, 150, 105))

    # 2. Vấn đề Mã nguồn & Khả năng tự sửa chữa
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("II. ĐÁNH GIÁ YÊU CẦU BÀN GIAO MÃ NGUỒN ĐỂ TỰ SỬA CHỮA KHI CÓ SỰ CỐ")
    format_run(r_h2, size_pt=12, bold=True, color_rgb=(3, 105, 161))

    p_mn = doc.add_paragraph()
    p_mn.paragraph_format.space_after = Pt(6)
    r_mn = p_mn.add_run(
        "Về mong muốn tự chủ kỹ thuật để khi có sự cố phần mềm thì cán bộ đơn vị có thể trực tiếp can thiệp, sửa chữa ngay, "
        "tôi xin phân tích rõ ranh giới kỹ thuật giữa hai hình thức bàn giao:"
    )
    format_run(r_mn, size_pt=11)

    # Bảng phân biệt bản deploy và source code
    tbl_mn = doc.add_table(rows=1, cols=3)
    tbl_mn.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mn.autofit = False

    headers_mn = ["Nội dung bàn giao", "Khả năng vận hành & Tự sửa lỗi", "Khả năng làm việc với đối tác"]
    col_w_mn = [Inches(1.8), Inches(3.2), Inches(2.0)]
    for i, h in enumerate(headers_mn):
        cell = tbl_mn.cell(0, i)
        cell.width = col_w_mn[i]
        set_cell_background(cell, "FEF3C7")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        format_run(r, size_pt=10.5, bold=True, color_rgb=(180, 83, 9))

    rows_mn_data = [
        (
            "Bản Triển Khai (Deploy/Publish)\n- File .dll đã biên dịch\n- Thư mục build Angular\n- File backup CSDL .bak",
            "- Vận hành: Chạy tốt 100% trên máy chủ HP cơ quan, cắt được 800k/tháng.\n- Sửa chữa: Sửa được cấu hình hệ thống (web.config) và sửa các câu lệnh Stored Procedure trong SQL Server.\n- Hạn chế: Không sửa được code logic C# gốc nếu chương trình bị lỗi thuật toán.",
            "RẤT DỄ ĐÒI HỎI BÀN GIAO:\nĐây là sản phẩm tối thiểu bên đối tác buộc phải cung cấp để đưa hệ thống lên máy chủ Chi nhánh."
        ),
        (
            "Toàn Bộ Mã Nguồn (Source Code)\n- Project Visual Studio C#\n- Toàn bộ source Angular gốc",
            "- Vận hành: Hoàn toàn chủ động 100%.\n- Sửa chữa: Lập trình viên có thể mở Visual Studio để sửa mọi dòng code, viết thêm chức năng mới tùy ý.",
            "CẦN THƯƠNG LƯỢNG:\nVới giá 800k/tháng, đối tác coi đây là tài sản trí tuệ. Cần kiểm tra lại hợp đồng ban đầu xem có điều khoản sở hữu mã nguồn hay không."
        )
    ]

    for row_idx, data in enumerate(rows_mn_data):
        row = tbl_mn.add_row()
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_w_mn[col_idx]
            set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            lines = text.split("\n")
            for line_idx, line in enumerate(lines):
                if line_idx == 0:
                    r = p.add_run(line + "\n")
                    format_run(r, size_pt=10, bold=True)
                else:
                    r = p.add_run(line + ("\n" if line_idx < len(lines)-1 else ""))
                    format_run(r, size_pt=9.5)

    for row in tbl_mn.rows:
        for cell in row.cells:
            tcPr = cell._element.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
            tcPr.append(tcBorders)

    # Lưu ý về nhân sự sửa chữa
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(8)
    p_note.paragraph_format.space_after = Pt(4)
    r_n = p_note.add_run(
        "Lưu ý quan trọng về năng lực can thiệp kỹ thuật: Phần mềm này viết bằng ngôn ngữ ASP.NET (C#) và CSDL SQL Server phức tạp. "
        "Để trực tiếp sửa chữa code khi có sự cố, đòi hỏi nhân sự phải có chuyên môn về Lập trình phần mềm (Software Developer). "
        "Cán bộ chuyên môn đất đai chỉ nên đảm nhận việc phân quyền tài khoản, cấu hình tham số báo cáo; "
        "còn các lỗi hệ thống sâu cần có nhân sự IT kỹ thuật chuyên trách để tránh nguy cơ làm gián đoạn dây chuyền tiếp nhận Một Cửa."
    )
    format_run(r_n, size_pt=10.5, italic=True)

    # 3. Kế hoạch di dời
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("III. KẾ HOẠCH TRIỂN KHAI 4 BƯỚC KHÔNG GIÁN ĐOẠN TIẾP NHẬN HỒ SƠ")
    format_run(r_h3, size_pt=12, bold=True, color_rgb=(3, 105, 161))

    steps = [
        ("Bước 1: Chuẩn bị máy chủ nội bộ (1-2 ngày)", "Cài đặt môi trường IIS, ASP.NET 4.8 và Microsoft SQL Server trên Máy chủ HP DL360 Gen9. Đăng ký mở IP tĩnh VNPT và mở cổng mạng an toàn trên Tường lửa FortiGate 61F."),
        ("Bước 2: Yêu cầu bàn giao dữ liệu & Chuyển quyền tên miền", "Làm việc với bên ông Trần Hồ Bảo Hoàng bàn giao: (1) File backup vpdkddtdm.bak; (2) Thư mục publish API; (3) Toàn bộ thư mục file scan đính kèm; (4) Chuyển nhượng quyền sở hữu tên miền vpdkddtdm.vn tại GMO-Z.com về Chi nhánh."),
        ("Bước 3: Dựng thử nghiệm (Staging) & Kiểm thử nội bộ", "Restore cơ sở dữ liệu lên máy chủ HP cơ quan, chạy thử nghiệm đầy đủ 9 phân hệ (Tiếp nhận -> Xử lý -> In biên nhận -> Trả kết quả) bảo đảm không có lỗi phát sinh."),
        ("Bước 4: Chuyển đổi chính thức vào cuối tuần", "Thực hiện trỏ bản ghi DNS tên miền vpdkddtdm.vn sang IP máy chủ Chi nhánh vào chiều thứ Sáu. Đến sáng thứ Hai, toàn bộ cán bộ đăng nhập và làm việc bình thường trên máy chủ mới của đơn vị, hoàn toàn không làm gián đoạn khâu tiếp nhận Một Cửa.")
    ]
    for s_title, s_desc in steps:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_after = Pt(4)
        r_st = p_s.add_run(s_title + ": ")
        format_run(r_st, size_pt=11, bold=True)
        r_sd = p_s.add_run(s_desc)
        format_run(r_sd, size_pt=10.5)

    add_signature_block(doc)
    doc.save(output_path)
    print(f"✅ Đã tạo xong Báo cáo 2: {output_path}")

if __name__ == "__main__":
    out1 = "/Users/admin/Desktop/BINHDUONG/Bao_cao_khao_sat_he_thong_OERP_vpdkddtdm.docx"
    out2 = "/Users/admin/Desktop/BINHDUONG/Bao_cao_phuong_an_di_doi_server_OERP_0_dong.docx"
    create_report_1(out1)
    create_report_2(out2)
