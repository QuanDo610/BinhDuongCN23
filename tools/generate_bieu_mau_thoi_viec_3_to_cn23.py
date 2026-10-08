import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borderless(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tcBorders>')
            tcPr.append(tcBorders)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", sz="4", color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    xml_str = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val != "none":
            xml_str += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        else:
            xml_str += f'<w:{side} w:val="none"/>'
    xml_str += '</w:tcBorders>'
    tcPr.append(parse_xml(xml_str))

def config_section_nd30(sec, is_first_section=False):
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.27)    # 210mm
    sec.page_height = Inches(11.69)  # 297mm
    # Can le chuan ND 30: Top 20mm, Bot 20mm, Left 30mm, Right 15mm
    sec.top_margin = Inches(0.79)    # 20mm
    sec.bottom_margin = Inches(0.79) # 20mm
    sec.left_margin = Inches(1.18)   # 30mm
    sec.right_margin = Inches(0.59)  # 15mm
    sec.different_first_page_header_footer = True
    
    # Restart page numbering at 1 for each independent document
    sectPr = sec._sectPr
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
    sectPr.append(pgNumType)
    
    header = sec.header
    if not is_first_section:
        header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    # Page number XML field
    fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="26"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>' % nsdecls('w'))
    p._p.append(fld)

def add_divider_line(container, length_type="nation"):
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 0.8
    dash_count = 18 if length_type == "nation" else 8
    r = p.add_run("—" * dash_count)
    r.font.name = "Times New Roman"
    r.font.size = Pt(8.5)
    r.font.bold = True
    return p

def create_document_for_team(team_name, team_short, tasks_description, output_path):
    doc = docx.Document()
    
    # Configure default Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)
    
    # =========================================================================
    # VĂN BẢN 1: PHIẾU TRÌNH (Tối ưu chuẩn 1 trang A4)
    # =========================================================================
    sec1 = doc.sections[0]
    config_section_nd30(sec1, is_first_section=True)
    
    # Header table: 2 cột không viền
    tbl_top = doc.add_table(rows=1, cols=2)
    tbl_top.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_top.autofit = False
    set_table_borderless(tbl_top)
    tbl_top.rows[0].cells[0].width = Inches(3.2)
    tbl_top.rows[0].cells[1].width = Inches(3.3)
    set_cell_margins(tbl_top.rows[0].cells[0], 0, 0, 0, 0)
    set_cell_margins(tbl_top.rows[0].cells[1], 0, 0, 0, 0)
    
    # Cột trái: Cơ quan ban hành
    c_left = tbl_top.rows[0].cells[0]
    p_unit = c_left.paragraphs[0]
    p_unit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_unit.paragraph_format.line_spacing = 1.15
    p_unit.paragraph_format.space_after = Pt(0)
    
    r1 = p_unit.add_run("VĂN PHÒNG ĐĂNG KÝ ĐẤT ĐAI\nTHÀNH PHỐ HỒ CHÍ MINH\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    
    r2 = p_unit.add_run("CHI NHÁNH SỐ 23\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.bold = True
    
    r3 = p_unit.add_run(f"{team_name.upper()}\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12.5)
    r3.font.bold = True
    
    add_divider_line(c_left, length_type="agency")
    
    # Cột phải: Quốc hiệu, tiêu ngữ & Ngày tháng
    c_right = tbl_top.rows[0].cells[1]
    p_nat = c_right.paragraphs[0]
    p_nat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nat.paragraph_format.line_spacing = 1.15
    p_nat.paragraph_format.space_after = Pt(0)
    
    r_nat1 = p_nat.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r_nat1.font.name = "Times New Roman"
    r_nat1.font.size = Pt(12.5)
    r_nat1.font.bold = True
    
    r_nat2 = p_nat.add_run("Độc lập - Tự do - Hạnh phúc")
    r_nat2.font.name = "Times New Roman"
    r_nat2.font.size = Pt(13)
    r_nat2.font.bold = True
    
    add_divider_line(c_right, length_type="nation")
    
    p_date = c_right.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(2)
    p_date.paragraph_format.space_after = Pt(0)
    r_d = p_date.add_run("Thành phố Hồ Chí Minh, ngày ..... tháng ..... năm 20.....")
    r_d.font.name = "Times New Roman"
    r_d.font.size = Pt(13)
    r_d.font.italic = True
    
    # Tên loại văn bản & Trích yếu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(1)
    r_tit = p_title.add_run("PHIẾU TRÌNH")
    r_tit.font.name = "Times New Roman"
    r_tit.font.size = Pt(14)
    r_tit.font.bold = True
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(1)
    r_sub = p_sub.add_run("Về việc giải quyết Đơn xin thôi việc của [ông/bà] ........................................")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    
    add_divider_line(doc, length_type="subject")
    
    # Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_kg.paragraph_format.first_line_indent = Inches(0.5)
    p_kg.paragraph_format.space_before = Pt(4)
    p_kg.paragraph_format.space_after = Pt(4)
    r_kg1 = p_kg.add_run("Kính gửi: ")
    r_kg1.font.name = "Times New Roman"
    r_kg1.font.size = Pt(13)
    r_kg1.font.bold = True
    r_kg2 = p_kg.add_run("Giám đốc Chi nhánh Văn phòng đăng ký đất đai số 23.")
    r_kg2.font.name = "Times New Roman"
    r_kg2.font.size = Pt(13)
    
    # Căn cứ
    cancu_list = [
        "Căn cứ Bộ luật Lao động ngày 20 tháng 11 năm 2019;",
        "Căn cứ Nghị định số 111/2022/NĐ-CP ngày 30 tháng 12 năm 2022 của Chính phủ về hợp đồng đối với một số loại công việc trong cơ quan hành chính và đơn vị sự nghiệp công lập;",
        "Căn cứ Quyết định số 2637/QĐ-SNNMT-VP ngày ..... tháng 12 năm 2025 của Sở Nông nghiệp và Môi trường Thành phố Hồ Chí Minh về việc ban hành Quy định chức năng, nhiệm vụ, quyền hạn và cơ cấu tổ chức của Chi nhánh Văn phòng đăng ký đất đai số 23 trực thuộc Văn phòng đăng ký đất đai Thành phố Hồ Chí Minh;"
    ]
    for cc in cancu_list:
        p_cc = doc.add_paragraph()
        p_cc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_cc.paragraph_format.first_line_indent = Inches(0.5)
        p_cc.paragraph_format.space_before = Pt(0)
        p_cc.paragraph_format.space_after = Pt(2)
        p_cc.paragraph_format.line_spacing = 1.15
        r_cc = p_cc.add_run(cc)
        r_cc.font.name = "Times New Roman"
        r_cc.font.size = Pt(13)
        r_cc.font.italic = True
    
    # Nội dung
    paras_pt = [
        f"Xét Đơn xin thôi việc ngày ..... tháng ..... năm 20..... của [ông/bà] .................................................... (sinh ngày: ...../...../.........., số CCCD: ...................................., ngày cấp: ...../...../.........., nơi cấp: ....................................). Hiện [ông/bà] .................................................... đang công tác tại {team_name}, đảm nhiệm nhiệm vụ: {tasks_description} và thực hiện một số công việc khác theo sự phân công, chỉ đạo của Lãnh đạo Tổ và Lãnh đạo Chi nhánh số 23.",
        f"Theo Biên bản cuộc họp ngày ..... tháng ..... năm 20..... của {team_name}, tập thể các thành phần dự họp đã tiến hành họp xem xét, thảo luận và thống nhất 100% về việc đề xuất giải quyết cho [ông/bà] .................................................... được thôi việc theo nguyện vọng cá nhân (kèm theo Biên bản cuộc họp).",
        f"Từ những nội dung trên, {team_name} kính báo cáo và xin ý kiến của Giám đốc Chi nhánh Văn phòng đăng ký đất đai số 23 xem xét, quyết định (kính chuyển Văn phòng đăng ký đất đai Thành phố Hồ Chí Minh xem xét ban hành Quyết định thôi việc và giải quyết các chế độ chính sách cho người lao động theo đúng quy định hiện hành)./."
    ]
    for p_text in paras_pt:
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.first_line_indent = Inches(0.5)
        p_body.paragraph_format.space_before = Pt(0)
        p_body.paragraph_format.space_after = Pt(4)
        p_body.paragraph_format.line_spacing = 1.15
        r_b = p_body.add_run(p_text)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)
    
    # Chữ ký Người trình & Tổ trưởng
    tbl_sign1 = doc.add_table(rows=1, cols=2)
    tbl_sign1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sign1.autofit = False
    set_table_borderless(tbl_sign1)
    tbl_sign1.rows[0].cells[0].width = Inches(3.2)
    tbl_sign1.rows[0].cells[1].width = Inches(3.3)
    
    # Cột trái: Người trình
    c_s1 = tbl_sign1.rows[0].cells[0]
    p_s1 = c_s1.paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s1.paragraph_format.line_spacing = 1.15
    p_s1.paragraph_format.space_after = Pt(0)
    r_s1_t = p_s1.add_run("NGƯỜI TRÌNH\n")
    r_s1_t.font.name = "Times New Roman"
    r_s1_t.font.size = Pt(13)
    r_s1_t.font.bold = True
    r_s1_g = p_s1.add_run("(Ký, ghi rõ họ tên)\n\n\n")
    r_s1_g.font.name = "Times New Roman"
    r_s1_g.font.size = Pt(12)
    r_s1_g.font.italic = True
    r_s1_n = p_s1.add_run("....................................................")
    r_s1_n.font.name = "Times New Roman"
    r_s1_n.font.size = Pt(13)
    
    # Cột phải: Tổ trưởng
    c_s2 = tbl_sign1.rows[0].cells[1]
    p_s2 = c_s2.paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s2.paragraph_format.line_spacing = 1.15
    p_s2.paragraph_format.space_after = Pt(0)
    r_s2_t = p_s2.add_run(f"TỔ TRƯỞNG\n{team_name.upper()}\n")
    r_s2_t.font.name = "Times New Roman"
    r_s2_t.font.size = Pt(13)
    r_s2_t.font.bold = True
    r_s2_g = p_s2.add_run("(Ký, ghi rõ họ tên)\n\n")
    r_s2_g.font.name = "Times New Roman"
    r_s2_g.font.size = Pt(12)
    r_s2_g.font.italic = True
    r_s2_n = p_s2.add_run("....................................................")
    r_s2_n.font.name = "Times New Roman"
    r_s2_n.font.size = Pt(13)
    
    # Khung Ý kiến của Giám đốc Chi nhánh số 23
    p_yk_title = doc.add_paragraph()
    p_yk_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_yk_title.paragraph_format.space_before = Pt(10)
    p_yk_title.paragraph_format.space_after = Pt(3)
    r_yk_tit = p_yk_title.add_run("Ý KIẾN CỦA GIÁM ĐỐC CHI NHÁNH VĂN PHÒNG ĐĂNG KÝ ĐẤT ĐAI SỐ 23:")
    r_yk_tit.font.name = "Times New Roman"
    r_yk_tit.font.size = Pt(13)
    r_yk_tit.font.bold = True
    
    # Tạo bảng khung chữ nhật cho ý kiến phê duyệt
    tbl_yk = doc.add_table(rows=1, cols=1)
    tbl_yk.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_yk.autofit = False
    cell_yk = tbl_yk.rows[0].cells[0]
    cell_yk.width = Inches(6.5)
    set_cell_borders(cell_yk, top="single", bottom="single", left="single", right="single", sz="4", color="888888")
    set_cell_margins(cell_yk, top=80, bottom=80, left=100, right=100)
    
    p_box = cell_yk.paragraphs[0]
    p_box.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_box.paragraph_format.space_after = Pt(3)
    r_box1 = p_box.add_run("................................................................................................................................................................................\n")
    r_box1.font.name = "Times New Roman"
    r_box1.font.size = Pt(13)
    r_box2 = p_box.add_run("................................................................................................................................................................................\n")
    r_box2.font.name = "Times New Roman"
    r_box2.font.size = Pt(13)
    r_box3 = p_box.add_run("................................................................................................................................................................................")
    r_box3.font.name = "Times New Roman"
    r_box3.font.size = Pt(13)
    
    p_gd_sign = cell_yk.add_paragraph()
    p_gd_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_gd_sign.paragraph_format.space_before = Pt(3)
    p_gd_sign.paragraph_format.space_after = Pt(0)
    p_gd_sign.paragraph_format.line_spacing = 1.15
    r_gd_d = p_gd_sign.add_run("Thành phố Hồ Chí Minh, ngày ..... tháng ..... năm 20.....\n")
    r_gd_d.font.name = "Times New Roman"
    r_gd_d.font.size = Pt(13)
    r_gd_d.font.italic = True
    r_gd_t = p_gd_sign.add_run("GIÁM ĐỐC                    \n")
    r_gd_t.font.name = "Times New Roman"
    r_gd_t.font.size = Pt(13)
    r_gd_t.font.bold = True
    r_gd_g = p_gd_sign.add_run("(Ký, ghi rõ họ tên, đóng dấu)      \n\n\n")
    r_gd_g.font.name = "Times New Roman"
    r_gd_g.font.size = Pt(12)
    r_gd_g.font.italic = True
    r_gd_name = p_gd_sign.add_run("....................................................        ")
    r_gd_name.font.name = "Times New Roman"
    r_gd_name.font.size = Pt(13)

    # =========================================================================
    # VĂN BẢN 2: BIÊN BẢN HỌP
    # =========================================================================
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    config_section_nd30(sec2, is_first_section=False)
    
    # Header table
    tbl_top2 = doc.add_table(rows=1, cols=2)
    tbl_top2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_top2.autofit = False
    set_table_borderless(tbl_top2)
    tbl_top2.rows[0].cells[0].width = Inches(3.2)
    tbl_top2.rows[0].cells[1].width = Inches(3.3)
    set_cell_margins(tbl_top2.rows[0].cells[0], 0, 0, 0, 0)
    set_cell_margins(tbl_top2.rows[0].cells[1], 0, 0, 0, 0)
    
    # Cột trái
    c2_left = tbl_top2.rows[0].cells[0]
    p2_unit = c2_left.paragraphs[0]
    p2_unit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_unit.paragraph_format.line_spacing = 1.15
    p2_unit.paragraph_format.space_after = Pt(0)
    
    r2_u1 = p2_unit.add_run("VĂN PHÒNG ĐĂNG KÝ ĐẤT ĐAI\nTHÀNH PHỐ HỒ CHÍ MINH\n")
    r2_u1.font.name = "Times New Roman"
    r2_u1.font.size = Pt(12)
    
    r2_u2 = p2_unit.add_run("CHI NHÁNH SỐ 23\n")
    r2_u2.font.name = "Times New Roman"
    r2_u2.font.size = Pt(12)
    r2_u2.font.bold = True
    
    r2_u3 = p2_unit.add_run(f"{team_name.upper()}\n")
    r2_u3.font.name = "Times New Roman"
    r2_u3.font.size = Pt(12.5)
    r2_u3.font.bold = True
    
    add_divider_line(c2_left, length_type="agency")
    
    # Cột phải
    c2_right = tbl_top2.rows[0].cells[1]
    p2_nat = c2_right.paragraphs[0]
    p2_nat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_nat.paragraph_format.line_spacing = 1.15
    p2_nat.paragraph_format.space_after = Pt(0)
    
    r2_nat1 = p2_nat.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r2_nat1.font.name = "Times New Roman"
    r2_nat1.font.size = Pt(12.5)
    r2_nat1.font.bold = True
    
    r2_nat2 = p2_nat.add_run("Độc lập - Tự do - Hạnh phúc")
    r2_nat2.font.name = "Times New Roman"
    r2_nat2.font.size = Pt(13)
    r2_nat2.font.bold = True
    
    add_divider_line(c2_right, length_type="nation")
    
    p2_date = c2_right.add_paragraph()
    p2_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_date.paragraph_format.space_before = Pt(2)
    p2_date.paragraph_format.space_after = Pt(0)
    r2_d = p2_date.add_run("Thành phố Hồ Chí Minh, ngày ..... tháng ..... năm 20.....")
    r2_d.font.name = "Times New Roman"
    r2_d.font.size = Pt(13)
    r2_d.font.italic = True
    
    # Tên biên bản
    p_bb_tit = doc.add_paragraph()
    p_bb_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_tit.paragraph_format.space_before = Pt(12)
    p_bb_tit.paragraph_format.space_after = Pt(1)
    r_bb = p_bb_tit.add_run("BIÊN BẢN")
    r_bb.font.name = "Times New Roman"
    r_bb.font.size = Pt(14)
    r_bb.font.bold = True
    
    p_bb_sub = doc.add_paragraph()
    p_bb_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_sub.paragraph_format.space_before = Pt(0)
    p_bb_sub.paragraph_format.space_after = Pt(1)
    p_bb_sub.paragraph_format.line_spacing = 1.15
    r_bbs1 = p_bb_sub.add_run(f"Họp về việc giải quyết Đơn xin thôi việc của [ông/bà] ........................................\nnhân viên thuộc {team_name} - Chi nhánh Văn phòng đăng ký đất đai số 23")
    r_bbs1.font.name = "Times New Roman"
    r_bbs1.font.size = Pt(13)
    r_bbs1.font.bold = True
    
    add_divider_line(doc, length_type="subject")
    
    # Thời gian địa điểm
    p_time = doc.add_paragraph()
    p_time.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_time.paragraph_format.first_line_indent = Inches(0.5)
    p_time.paragraph_format.space_before = Pt(4)
    p_time.paragraph_format.space_after = Pt(4)
    r_tm = p_time.add_run(f"Vào lúc ..... giờ ..... phút, ngày ..... tháng ..... năm 20....., tại Trụ sở Chi nhánh Văn phòng đăng ký đất đai số 23 (địa chỉ: Số 358 đường Huỳnh Văn Cù, phường Thủ Dầu Một, Thành phố Hồ Chí Minh), {team_name} tổ chức cuộc họp về việc xem xét giải quyết Đơn xin thôi việc của [ông/bà] .................................................... với thành phần và nội dung cụ thể như sau:")
    r_tm.font.name = "Times New Roman"
    r_tm.font.size = Pt(13)
    
    # I. Thành phần cuộc họp
    p_tp = doc.add_paragraph()
    p_tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_tp.paragraph_format.space_before = Pt(4)
    p_tp.paragraph_format.space_after = Pt(2)
    r_tp = p_tp.add_run("I. Thành phần cuộc họp:")
    r_tp.font.name = "Times New Roman"
    r_tp.font.size = Pt(13)
    r_tp.font.bold = True
    
    tp_items = [
        "- [Ông/Bà] .................................................... - Tổ trưởng - Chủ trì;",
        "- [Ông/Bà] .................................................... - Tổ phó;",
        "- [Ông/Bà] .................................................... - Đại diện Tổ Công đoàn;",
        "- [Ông/Bà] .................................................... - Thư ký cuộc họp;",
        "- Người có đơn xin thôi việc: [Ông/Bà] ....................................................;",
        f"- Cùng toàn thể các viên chức, người lao động thuộc {team_name}:",
        "  + [Ông/Bà] .................................................... - Vị trí: ....................................................;",
        "  + [Ông/Bà] .................................................... - Vị trí: ....................................................;",
        "  (Tổng số có mặt: ...../..... đồng chí; Vắng mặt: ..... đồng chí, lý do: ........................................)."
    ]
    for item in tp_items:
        p_it = doc.add_paragraph()
        p_it.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_it.paragraph_format.first_line_indent = Inches(0.5)
        p_it.paragraph_format.space_before = Pt(0)
        p_it.paragraph_format.space_after = Pt(2)
        p_it.paragraph_format.line_spacing = 1.15
        r_it = p_it.add_run(item)
        r_it.font.name = "Times New Roman"
        r_it.font.size = Pt(13)
    
    # II. Nội dung cuộc họp
    p_nd = doc.add_paragraph()
    p_nd.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_nd.paragraph_format.space_before = Pt(4)
    p_nd.paragraph_format.space_after = Pt(2)
    r_nd = p_nd.add_run("II. Nội dung cuộc họp:")
    r_nd.font.name = "Times New Roman"
    r_nd.font.size = Pt(13)
    r_nd.font.bold = True
    
    cc_bb = [
        "- Căn cứ Bộ luật Lao động ngày 20 tháng 11 năm 2019;",
        "- Căn cứ Nghị định số 111/2022/NĐ-CP ngày 30 tháng 12 năm 2022 của Chính phủ về hợp đồng đối với một số loại công việc trong cơ quan hành chính và đơn vị sự nghiệp công lập;",
        "- Căn cứ Quyết định số 2637/QĐ-SNNMT-VP ngày ..... tháng 12 năm 2025 của Sở Nông nghiệp và Môi trường Thành phố Hồ Chí Minh về việc ban hành Quy định chức năng, nhiệm vụ, quyền hạn và cơ cấu tổ chức của Chi nhánh Văn phòng đăng ký đất đai số 23;",
        "- Căn cứ Đơn xin thôi việc ngày ..... tháng ..... năm 20..... của [ông/bà] ....................................................;"
    ]
    for c_item in cc_bb:
        p_cbb = doc.add_paragraph()
        p_cbb.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_cbb.paragraph_format.first_line_indent = Inches(0.5)
        p_cbb.paragraph_format.space_before = Pt(0)
        p_cbb.paragraph_format.space_after = Pt(2)
        p_cbb.paragraph_format.line_spacing = 1.15
        r_cbb = p_cbb.add_run(c_item)
        r_cbb.font.name = "Times New Roman"
        r_cbb.font.size = Pt(13)
        r_cbb.font.italic = True
    
    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    p_intro.paragraph_format.space_before = Pt(3)
    p_intro.paragraph_format.space_after = Pt(3)
    r_in = p_intro.add_run("Sau khi nghe Thư ký cuộc họp thông qua nội dung chương trình họp, các thành phần dự họp tiến hành thảo luận và thống nhất các nội dung sau:")
    r_in.font.name = "Times New Roman"
    r_in.font.size = Pt(13)
    
    # Mục 1: Ý kiến của người có đơn
    p_yk_ng = doc.add_paragraph()
    p_yk_ng.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_yk_ng.paragraph_format.first_line_indent = Inches(0.5)
    p_yk_ng.paragraph_format.space_before = Pt(2)
    p_yk_ng.paragraph_format.space_after = Pt(3)
    r_ykn1 = p_yk_ng.add_run("1. Ý kiến của [ông/bà] .................................................... (người xin thôi việc): ")
    r_ykn1.font.name = "Times New Roman"
    r_ykn1.font.size = Pt(13)
    r_ykn1.font.bold = True
    r_ykn2 = p_yk_ng.add_run("Do điều kiện hoàn cảnh gia đình (hoặc lý do riêng chính đáng) nên tôi không thể tiếp tục đảm nhiệm công việc hiện tại. Kính mong Lãnh đạo Chi nhánh số 23 và Lãnh đạo Tổ tạo điều kiện cho phép tôi được nghỉ việc kể từ ngày ...../...../20..... Bản thân cam kết sẽ chấp hành nghiêm túc nội quy cơ quan, hoàn thành tốt các nhiệm vụ được giao cho đến ngày nghỉ và thực hiện bàn giao đầy đủ toàn bộ hồ sơ, tài liệu, tài sản, trang thiết bị máy móc cũng như các công việc đang phụ trách theo đúng quy định.")
    r_ykn2.font.name = "Times New Roman"
    r_ykn2.font.size = Pt(13)
    
    # Mục 2: Ý kiến thành phần dự họp
    p_yk_tp = doc.add_paragraph()
    p_yk_tp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_yk_tp.paragraph_format.first_line_indent = Inches(0.5)
    p_yk_tp.paragraph_format.space_before = Pt(2)
    p_yk_tp.paragraph_format.space_after = Pt(3)
    r_yktp1 = p_yk_tp.add_run("2. Ý kiến của các thành phần dự họp:\n")
    r_yktp1.font.name = "Times New Roman"
    r_yktp1.font.size = Pt(13)
    r_yktp1.font.bold = True
    r_yktp2 = p_yk_tp.add_run(f"- Ý kiến của các đồng chí trong {team_name}: Ghi nhận tinh thần làm việc, trách nhiệm và kết quả đóng góp của [ông/bà] .................................................... trong thời gian công tác. Hoàn toàn đồng thuận và thống nhất theo nguyện vọng xin thôi việc của [ông/bà] ....................................................\n"
                             "- Ý kiến của đại diện Tổ Công đoàn: Thống nhất theo nguyện vọng xin thôi việc của đoàn viên. Đề nghị Lãnh đạo Tổ và Lãnh đạo Chi nhánh xem xét giải quyết đầy đủ quyền lợi, chế độ chính sách thôi việc cho người lao động theo đúng quy định của pháp luật lao động.")
    r_yktp2.font.name = "Times New Roman"
    r_yktp2.font.size = Pt(13)
    
    # Mục 3: Kết luận của Chủ trì
    p_kl = doc.add_paragraph()
    p_kl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kl.paragraph_format.first_line_indent = Inches(0.5)
    p_kl.paragraph_format.space_before = Pt(2)
    p_kl.paragraph_format.space_after = Pt(3)
    r_kl1 = p_kl.add_run("3. Kết luận của Chủ trì cuộc họp (Tổ trưởng):\n")
    r_kl1.font.name = "Times New Roman"
    r_kl1.font.size = Pt(13)
    r_kl1.font.bold = True
    r_kl2 = p_kl.add_run(f"- Tập thể {team_name} thống nhất 100% (...../..... thành viên có mặt) đồng ý đề xuất giải quyết cho [ông/bà] .................................................... được thôi việc kể từ ngày ..... tháng ..... năm 20..... theo nguyện vọng cá nhân.\n"
                         "- Về công tác bàn giao: Phân công [ông/bà] .................................................... chịu trách nhiệm tiếp nhận bàn giao toàn bộ hồ sơ nghiệp vụ, cơ sở dữ liệu, sổ sách, tài sản trang thiết bị và các công việc do [ông/bà] .................................................... đang quản lý, phụ trách. Yêu cầu hoàn thành việc bàn giao trước ngày ...../...../20..... (lập Biên bản bàn giao kèm theo).\n"
                         "- Giao Thư ký cuộc họp phối hợp hoàn thiện bộ hồ sơ (gồm: Phiếu trình, Biên bản cuộc họp, Đơn xin thôi việc) kính trình Giám đốc Chi nhánh Văn phòng đăng ký đất đai số 23 xem xét, báo cáo Ban Giám đốc Văn phòng đăng ký đất đai Thành phố Hồ Chí Minh giải quyết chế độ nghỉ việc theo đúng quy định cho [ông/bà] ....................................................")
    r_kl2.font.name = "Times New Roman"
    r_kl2.font.size = Pt(13)
    
    # Kết thúc
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_end.paragraph_format.first_line_indent = Inches(0.5)
    p_end.paragraph_format.space_before = Pt(3)
    p_end.paragraph_format.space_after = Pt(6)
    r_end = p_end.add_run("Biên bản được thống nhất và thông qua lúc ..... giờ ..... phút cùng ngày./.")
    r_end.font.name = "Times New Roman"
    r_end.font.size = Pt(13)
    r_end.font.italic = True
    
    # Bảng ký Biên bản: Hàng 1 (Chủ trì, Tổ phó, Công đoàn)
    tbl_bb_sign1 = doc.add_table(rows=1, cols=3)
    tbl_bb_sign1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_bb_sign1.autofit = False
    set_table_borderless(tbl_bb_sign1)
    tbl_bb_sign1.rows[0].cells[0].width = Inches(2.2)
    tbl_bb_sign1.rows[0].cells[1].width = Inches(2.1)
    tbl_bb_sign1.rows[0].cells[2].width = Inches(2.2)
    
    # Chủ trì
    c_bb_ct = tbl_bb_sign1.rows[0].cells[0]
    p_bb_ct = c_bb_ct.paragraphs[0]
    p_bb_ct.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_ct.paragraph_format.line_spacing = 1.15
    p_bb_ct.paragraph_format.space_after = Pt(0)
    p_bb_ct.add_run("CHỦ TRÌ\n").font.bold = True
    p_bb_ct.runs[0].font.size = Pt(13)
    p_bb_ct.runs[0].font.name = "Times New Roman"
    r_g1 = p_bb_ct.add_run("Tổ trưởng\n(Ký, ghi rõ họ tên)\n\n\n")
    r_g1.font.size = Pt(12)
    r_g1.font.italic = True
    r_g1.font.name = "Times New Roman"
    r_n1 = p_bb_ct.add_run("....................................................")
    r_n1.font.size = Pt(13)
    r_n1.font.name = "Times New Roman"
    
    # Tổ phó
    c_bb_tp = tbl_bb_sign1.rows[0].cells[1]
    p_bb_tp = c_bb_tp.paragraphs[0]
    p_bb_tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_tp.paragraph_format.line_spacing = 1.15
    p_bb_tp.paragraph_format.space_after = Pt(0)
    p_bb_tp.add_run("TỔ PHÓ\n").font.bold = True
    p_bb_tp.runs[0].font.size = Pt(13)
    p_bb_tp.runs[0].font.name = "Times New Roman"
    r_g2 = p_bb_tp.add_run("\n(Ký, ghi rõ họ tên)\n\n\n")
    r_g2.font.size = Pt(12)
    r_g2.font.italic = True
    r_g2.font.name = "Times New Roman"
    r_n2 = p_bb_tp.add_run("....................................................")
    r_n2.font.size = Pt(13)
    r_n2.font.name = "Times New Roman"
    
    # Đại diện Công đoàn
    c_bb_cd = tbl_bb_sign1.rows[0].cells[2]
    p_bb_cd = c_bb_cd.paragraphs[0]
    p_bb_cd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_cd.paragraph_format.line_spacing = 1.15
    p_bb_cd.paragraph_format.space_after = Pt(0)
    p_bb_cd.add_run("ĐẠI DIỆN CÔNG ĐOÀN\n").font.bold = True
    p_bb_cd.runs[0].font.size = Pt(13)
    p_bb_cd.runs[0].font.name = "Times New Roman"
    r_g3 = p_bb_cd.add_run("\n(Ký, ghi rõ họ tên)\n\n\n")
    r_g3.font.size = Pt(12)
    r_g3.font.italic = True
    r_g3.font.name = "Times New Roman"
    r_n3 = p_bb_cd.add_run("....................................................")
    r_n3.font.size = Pt(13)
    r_n3.font.name = "Times New Roman"
    
    # Bảng ký Hàng 2: Thư ký & Người xin thôi việc
    tbl_bb_sign2 = doc.add_table(rows=1, cols=2)
    tbl_bb_sign2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_bb_sign2.autofit = False
    set_table_borderless(tbl_bb_sign2)
    tbl_bb_sign2.rows[0].cells[0].width = Inches(3.2)
    tbl_bb_sign2.rows[0].cells[1].width = Inches(3.3)
    
    # Thư ký
    c_bb_tk = tbl_bb_sign2.rows[0].cells[0]
    p_bb_tk = c_bb_tk.paragraphs[0]
    p_bb_tk.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_tk.paragraph_format.line_spacing = 1.15
    p_bb_tk.paragraph_format.space_after = Pt(0)
    p_bb_tk.add_run("THƯ KÝ CUỘC HỌP\n").font.bold = True
    p_bb_tk.runs[0].font.size = Pt(13)
    p_bb_tk.runs[0].font.name = "Times New Roman"
    r_g4 = p_bb_tk.add_run("(Ký, ghi rõ họ tên)\n\n\n")
    r_g4.font.size = Pt(12)
    r_g4.font.italic = True
    r_g4.font.name = "Times New Roman"
    r_n4 = p_bb_tk.add_run("....................................................")
    r_n4.font.size = Pt(13)
    r_n4.font.name = "Times New Roman"
    
    # Người xin thôi việc
    c_bb_nv = tbl_bb_sign2.rows[0].cells[1]
    p_bb_nv = c_bb_nv.paragraphs[0]
    p_bb_nv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bb_nv.paragraph_format.line_spacing = 1.15
    p_bb_nv.paragraph_format.space_after = Pt(0)
    p_bb_nv.add_run("NGƯỜI XIN THÔI VIỆC\n").font.bold = True
    p_bb_nv.runs[0].font.size = Pt(13)
    p_bb_nv.runs[0].font.name = "Times New Roman"
    r_g5 = p_bb_nv.add_run("(Ký, ghi rõ họ tên)\n\n\n")
    r_g5.font.size = Pt(12)
    r_g5.font.italic = True
    r_g5.font.name = "Times New Roman"
    r_n5 = p_bb_nv.add_run("....................................................")
    r_n5.font.size = Pt(13)
    r_n5.font.name = "Times New Roman"

    # =========================================================================
    # VĂN BẢN 3: ĐƠN XIN THÔI VIỆC (Tối ưu chuẩn 1 trang A4)
    # =========================================================================
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    config_section_nd30(sec3, is_first_section=False)
    
    # Quốc hiệu, tiêu ngữ (Đơn xin thôi việc của cá nhân canh giữa trang)
    p3_nat = doc.add_paragraph()
    p3_nat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3_nat.paragraph_format.line_spacing = 1.15
    p3_nat.paragraph_format.space_after = Pt(0)
    
    r3_nat1 = p3_nat.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r3_nat1.font.name = "Times New Roman"
    r3_nat1.font.size = Pt(13)
    r3_nat1.font.bold = True
    
    r3_nat2 = p3_nat.add_run("Độc lập - Tự do - Hạnh phúc")
    r3_nat2.font.name = "Times New Roman"
    r3_nat2.font.size = Pt(13.5)
    r3_nat2.font.bold = True
    
    add_divider_line(doc, length_type="nation")
    
    # Tiêu đề Đơn
    p_don_tit = doc.add_paragraph()
    p_don_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_don_tit.paragraph_format.space_before = Pt(10)
    p_don_tit.paragraph_format.space_after = Pt(2)
    r_don = p_don_tit.add_run("ĐƠN XIN THÔI VIỆC")
    r_don.font.name = "Times New Roman"
    r_don.font.size = Pt(15)
    r_don.font.bold = True
    
    p_don_sub = doc.add_paragraph()
    p_don_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_don_sub.paragraph_format.space_before = Pt(0)
    p_don_sub.paragraph_format.space_after = Pt(8)
    r_dsub = p_don_sub.add_run("(V/v xin nghỉ việc theo nguyện vọng cá nhân)")
    r_dsub.font.name = "Times New Roman"
    r_dsub.font.size = Pt(12)
    r_dsub.font.italic = True
    
    # Kính gửi
    p_don_kg = doc.add_paragraph()
    p_don_kg.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_don_kg.paragraph_format.first_line_indent = Inches(0.5)
    p_don_kg.paragraph_format.space_before = Pt(2)
    p_don_kg.paragraph_format.space_after = Pt(2)
    r_dkg1 = p_don_kg.add_run("Kính gửi:\n")
    r_dkg1.font.name = "Times New Roman"
    r_dkg1.font.size = Pt(13)
    r_dkg1.font.bold = True
    
    kg_targets = [
        "- Ban Giám đốc Văn phòng Đăng ký đất đai Thành phố Hồ Chí Minh;",
        "- Ban Giám đốc Chi nhánh Văn phòng đăng ký đất đai số 23;",
        f"- Lãnh đạo {team_name}."
    ]
    for kg_t in kg_targets:
        p_kgt = doc.add_paragraph()
        p_kgt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_kgt.paragraph_format.first_line_indent = Inches(1.0)
        p_kgt.paragraph_format.space_before = Pt(0)
        p_kgt.paragraph_format.space_after = Pt(2)
        r_kgt = p_kgt.add_run(kg_t)
        r_kgt.font.name = "Times New Roman"
        r_kgt.font.size = Pt(13)
    
    # Thông tin người làm đơn
    info_fields = [
        ("Tôi tên là:", "..........................................................................................", "Giới tính:", "...................."),
        ("Ngày, tháng, năm sinh:", "...../...../..........", "Nơi sinh:", ".................................................."),
        ("CCCD/Số định danh:", "..................................................", "Cấp ngày:", "...../...../.........."),
        ("Nơi cấp CCCD:", "....................................................................................................................", "", ""),
        ("Nơi đăng ký thường trú:", ".................................................................................................................", "", ""),
        ("Chỗ ở hiện nay:", "........................................................................................................................", "", ""),
        ("Điện thoại liên hệ:", "..................................................", "Email:", ".................................................."),
        ("Vị trí việc làm/Chức danh:", "................................................................................................................", "", ""),
        ("Bộ phận công tác:", f"{team_name}, Chi nhánh Văn phòng đăng ký đất đai số 23.", "", ""),
        ("Hợp đồng lao động số:", "..................................................", "Ký ngày:", "...../...../..........")
    ]
    for f in info_fields:
        p_f = doc.add_paragraph()
        p_f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_f.paragraph_format.first_line_indent = Inches(0.5)
        p_f.paragraph_format.space_before = Pt(0)
        p_f.paragraph_format.space_after = Pt(2)
        p_f.paragraph_format.line_spacing = 1.15
        
        r_lbl1 = p_f.add_run(f[0] + " ")
        r_lbl1.font.name = "Times New Roman"
        r_lbl1.font.size = Pt(13)
        r_lbl1.font.bold = True
        
        r_val1 = p_f.add_run(f[1] + "  ")
        r_val1.font.name = "Times New Roman"
        r_val1.font.size = Pt(13)
        
        if f[2]:
            r_lbl2 = p_f.add_run(f[2] + " ")
            r_lbl2.font.name = "Times New Roman"
            r_lbl2.font.size = Pt(13)
            r_lbl2.font.bold = True
            
            r_val2 = p_f.add_run(f[3])
            r_val2.font.name = "Times New Roman"
            r_val2.font.size = Pt(13)
    
    # Nội dung đơn
    don_texts = [
        f"Nay tôi làm đơn này kính xin Ban Giám đốc Văn phòng Đăng ký đất đai Thành phố Hồ Chí Minh, Ban Giám đốc Chi nhánh Văn phòng đăng ký đất đai số 23 và Lãnh đạo {team_name} xem xét cho tôi được thôi việc kể từ ngày ..... tháng ..... năm 20.....",
        "Lý do xin thôi việc: ............................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................................",
        "Tôi cam kết sẽ nghiêm túc chấp hành mọi quy định, tiếp tục hoàn thành tốt các công việc được giao cho đến ngày chính thức nghỉ việc; đồng thời hoàn tất đầy đủ thủ tục bàn giao toàn bộ hồ sơ, tài liệu, số liệu chuyên môn, tài sản, trang thiết bị máy móc và các công việc đang đảm nhiệm cho người được phân công tiếp nhận theo đúng quy định của đơn vị.",
        "Kính mong Ban Giám đốc Văn phòng Đăng ký đất đai Thành phố, Ban Giám đốc Chi nhánh số 23 và Lãnh đạo Tổ quan tâm xem xét, tạo điều kiện chấp thuận giải quyết cho tôi được thôi việc theo nguyện vọng cá nhân.",
        "Tôi xin chân thành cảm ơn!./."
    ]
    for d_txt in don_texts:
        p_dbody = doc.add_paragraph()
        p_dbody.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_dbody.paragraph_format.first_line_indent = Inches(0.5)
        p_dbody.paragraph_format.space_before = Pt(1)
        p_dbody.paragraph_format.space_after = Pt(3)
        p_dbody.paragraph_format.line_spacing = 1.15
        r_db = p_dbody.add_run(d_txt)
        r_db.font.name = "Times New Roman"
        r_db.font.size = Pt(13)
    
    # Ngày tháng & Người viết đơn
    p_nv_sign = doc.add_paragraph()
    p_nv_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_nv_sign.paragraph_format.space_before = Pt(4)
    p_nv_sign.paragraph_format.space_after = Pt(0)
    p_nv_sign.paragraph_format.line_spacing = 1.15
    r_nvd = p_nv_sign.add_run("Thành phố Hồ Chí Minh, ngày ..... tháng ..... năm 20.....\n")
    r_nvd.font.name = "Times New Roman"
    r_nvd.font.size = Pt(13)
    r_nvd.font.italic = True
    r_nvt = p_nv_sign.add_run("NGƯỜI VIẾT ĐƠN                    \n")
    r_nvt.font.name = "Times New Roman"
    r_nvt.font.size = Pt(13)
    r_nvt.font.bold = True
    r_nvg = p_nv_sign.add_run("(Ký, ghi rõ họ tên)               \n\n\n")
    r_nvg.font.name = "Times New Roman"
    r_nvg.font.size = Pt(12)
    r_nvg.font.italic = True
    r_nvname = p_nv_sign.add_run("....................................................        ")
    r_nvname.font.name = "Times New Roman"
    r_nvname.font.size = Pt(13)
    
    # Ý kiến xác nhận: Bảng 2 cột (Tổ trưởng & Giám đốc Chi nhánh)
    tbl_don_yk = doc.add_table(rows=1, cols=2)
    tbl_don_yk.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_don_yk.autofit = False
    tbl_don_yk.rows[0].cells[0].width = Inches(3.2)
    tbl_don_yk.rows[0].cells[1].width = Inches(3.3)
    
    # Khung viền nhẹ phân tách
    for c_i in tbl_don_yk.rows[0].cells:
        set_cell_borders(c_i, top="single", bottom="single", left="single", right="single", sz="4", color="888888")
        set_cell_margins(c_i, top=60, bottom=60, left=80, right=80)
    
    # Cột trái: Ý kiến của Tổ
    c_yk_to = tbl_don_yk.rows[0].cells[0]
    p_ykt = c_yk_to.paragraphs[0]
    p_ykt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ykt.paragraph_format.line_spacing = 1.15
    p_ykt.paragraph_format.space_after = Pt(2)
    r_ykt_tit = p_ykt.add_run(f"Ý KIẾN CỦA {team_name.upper()}\n")
    r_ykt_tit.font.name = "Times New Roman"
    r_ykt_tit.font.size = Pt(12)
    r_ykt_tit.font.bold = True
    
    p_ykt_body = c_yk_to.add_paragraph()
    p_ykt_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ykt_body.paragraph_format.line_spacing = 1.15
    p_ykt_body.paragraph_format.space_after = Pt(2)
    r_ykt_con = p_ykt_body.add_run("Kính trình Ban Giám đốc Chi nhánh số 23 xem xét, chấp thuận nguyện vọng xin thôi việc của [ông/bà] ....................................................\n"
                                  "................................................................................")
    r_ykt_con.font.name = "Times New Roman"
    r_ykt_con.font.size = Pt(12)
    r_ykt_con.font.italic = True
    
    p_ykt_sig = c_yk_to.add_paragraph()
    p_ykt_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ykt_sig.paragraph_format.line_spacing = 1.15
    p_ykt_sig.paragraph_format.space_after = Pt(0)
    p_ykt_sig.paragraph_format.space_before = Pt(3)
    r_ykt_date = p_ykt_sig.add_run("Ngày ..... tháng ..... năm 20.....\n")
    r_ykt_date.font.size = Pt(12)
    r_ykt_date.font.italic = True
    r_ykt_date.font.name = "Times New Roman"
    r_ykt_post = p_ykt_sig.add_run("TỔ TRƯỞNG\n")
    r_ykt_post.font.size = Pt(12.5)
    r_ykt_post.font.bold = True
    r_ykt_post.font.name = "Times New Roman"
    r_ykt_not = p_ykt_sig.add_run("(Ký, ghi rõ họ tên)\n\n\n")
    r_ykt_not.font.size = Pt(11)
    r_ykt_not.font.italic = True
    r_ykt_not.font.name = "Times New Roman"
    r_ykt_name = p_ykt_sig.add_run("....................................................")
    r_ykt_name.font.size = Pt(12)
    r_ykt_name.font.name = "Times New Roman"
    
    # Cột phải: Ý kiến Ban Giám đốc Chi nhánh số 23
    c_yk_bgd = tbl_don_yk.rows[0].cells[1]
    p_ykbgd = c_yk_bgd.paragraphs[0]
    p_ykbgd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ykbgd.paragraph_format.line_spacing = 1.15
    p_ykbgd.paragraph_format.space_after = Pt(2)
    r_ykb_tit = p_ykbgd.add_run("Ý KIẾN BAN GIÁM ĐỐC\nCHI NHÁNH SỐ 23\n")
    r_ykb_tit.font.name = "Times New Roman"
    r_ykb_tit.font.size = Pt(12)
    r_ykb_tit.font.bold = True
    
    p_ykb_body = c_yk_bgd.add_paragraph()
    p_ykb_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ykb_body.paragraph_format.line_spacing = 1.15
    p_ykb_body.paragraph_format.space_after = Pt(2)
    r_ykb_con = p_ykb_body.add_run("Kính chuyển Văn phòng đăng ký đất đai Thành phố Hồ Chí Minh xem xét, giải quyết chế độ thôi việc theo quy định.\n"
                                  "................................................................................")
    r_ykb_con.font.name = "Times New Roman"
    r_ykb_con.font.size = Pt(12)
    r_ykb_con.font.italic = True
    
    p_ykb_sig = c_yk_bgd.add_paragraph()
    p_ykb_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ykb_sig.paragraph_format.line_spacing = 1.15
    p_ykb_sig.paragraph_format.space_after = Pt(0)
    p_ykb_sig.paragraph_format.space_before = Pt(3)
    r_ykb_date = p_ykb_sig.add_run("Ngày ..... tháng ..... năm 20.....\n")
    r_ykb_date.font.size = Pt(12)
    r_ykb_date.font.italic = True
    r_ykb_date.font.name = "Times New Roman"
    r_ykb_post = p_ykb_sig.add_run("GIÁM ĐỐC\n")
    r_ykb_post.font.size = Pt(12.5)
    r_ykb_post.font.bold = True
    r_ykb_post.font.name = "Times New Roman"
    r_ykb_not = p_ykb_sig.add_run("(Ký, ghi rõ họ tên, đóng dấu)\n\n\n")
    r_ykb_not.font.size = Pt(11)
    r_ykb_not.font.italic = True
    r_ykb_not.font.name = "Times New Roman"
    r_ykb_name = p_ykb_sig.add_run("....................................................")
    r_ykb_name.font.size = Pt(12)
    r_ykb_name.font.name = "Times New Roman"
    
    doc.save(output_path)
    print(f"Generated: {output_path}")

def main():
    teams = [
        {
            "name": "Tổ Hành chính – Tổng hợp",
            "short": "Tổ HC-TH",
            "tasks": "công tác văn thư, lưu trữ nội bộ, kế toán, thủ quỹ, tiếp nhận hồ sơ và trả kết quả thủ tục hành chính, công nghệ thông tin và công tác tổng hợp hành chính",
            "file": "OUTPUT/Mau_ho_so_thoi_viec_To_Hanh_chinh_Tong_hop_CN23.docx"
        },
        {
            "name": "Tổ Đăng ký và Cấp giấy chứng nhận",
            "short": "Tổ ĐK-CGCN",
            "tasks": "kiểm tra, thẩm tra hồ sơ đăng ký đất đai, đăng ký biến động quyền sử dụng đất, tài sản gắn liền với đất, cấp và đính chính Giấy chứng nhận, đăng ký biện pháp bảo đảm",
            "file": "OUTPUT/Mau_ho_so_thoi_viec_To_Dang_ky_va_Cap_GCN_CN23.docx"
        },
        {
            "name": "Tổ Kỹ thuật Địa chính và Lưu trữ",
            "short": "Tổ KTĐC-LT",
            "tasks": "đo đạc, lập và chỉnh lý bản đồ địa chính, trích lục, trích đo thửa đất, cập nhật cơ sở dữ liệu địa chính, số hóa và quản lý lưu trữ hồ sơ địa chính",
            "file": "OUTPUT/Mau_ho_so_thoi_viec_To_Ky_thuat_Dia_chinh_va_Luu_tru_CN23.docx"
        }
    ]
    
    os.makedirs("OUTPUT", exist_ok=True)
    for t in teams:
        create_document_for_team(t["name"], t["short"], t["tasks"], t["file"])

if __name__ == "__main__":
    main()
