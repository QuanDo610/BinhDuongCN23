import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_master(file_path):
    wb = openpyxl.Workbook()
    
    # --- 1. Tạo Sheet MUC TIEU ---
    ws_muc_tieu = wb.active
    ws_muc_tieu.title = "MUC TIEU"
    
    headers_mt = ["STT", "Họ và Tên", "Phường/Xã", "Nhiệm vụ", "Tổng khối lượng giao (Thửa/Hồ sơ)"]
    ws_muc_tieu.append(headers_mt)
    
    # Mock data MUC TIEU
    data_mt = [
        [1, "Nguyễn Thị Thùy Trang", "Phường Thủ Dầu Một", "Rà soát", 2126],
        [2, "Nguyễn Trung Thành", "Phường Thủ Dầu Một", "Rà soát", 2126],
        [3, "Huỳnh Thành Duy", "Kho lưu trữ", "Lục hồ sơ", 4371],
        [4, "Nguyễn Nhật Quyền Trinh", "Kho lưu trữ", "Scan hồ sơ", 4371],
        [5, "Nguyễn Anh Khôi", "Bình Dương", "Không gian", 13620],
    ]
    for r in data_mt:
        ws_muc_tieu.append(r)
        
    # --- 2. Tạo Sheet NHAT KY ---
    ws_nhat_ky = wb.create_sheet(title="NHAT KY")
    headers_nk = ["Ngày thực hiện", "Họ và Tên", "Loại công việc", "Chi tiết công việc", "Khối lượng hoàn thành (Số lượng)", "Ghi chú"]
    ws_nhat_ky.append(headers_nk)
    
    # Data Validation for NHAT KY
    dv_name = DataValidation(type="list", formula1='="Nguyễn Thị Thùy Trang,Nguyễn Trung Thành,Huỳnh Thành Duy,Nguyễn Nhật Quyền Trinh,Nguyễn Anh Khôi"', allow_blank=True)
    dv_task = DataValidation(type="list", formula1='="Rà soát,Lục hồ sơ,Scan hồ sơ,Không gian"', allow_blank=True)
    dv_detail = DataValidation(type="list", formula1='="Rà gửi lục,Cập nhật lên nhóm 1,Đã lục giao lên scan,Đã scan từ kho lục lên,Chỉnh lý thiếu không gian"', allow_blank=True)
    
    # Add validation to columns (from row 2 to 1000)
    dv_name.add('B2:B1000')
    dv_task.add('C2:C1000')
    dv_detail.add('D2:D1000')
    
    ws_nhat_ky.add_data_validation(dv_name)
    ws_nhat_ky.add_data_validation(dv_task)
    ws_nhat_ky.add_data_validation(dv_detail)
    
    # Mock data NHAT KY
    data_nk = [
        ["10/08/2026", "Nguyễn Thị Thùy Trang", "Rà soát", "Rà gửi lục", 108, "Thử nghiệm"],
        ["10/08/2026", "Nguyễn Thị Thùy Trang", "Rà soát", "Cập nhật lên nhóm 1", 46, "Thử nghiệm"],
        ["10/08/2026", "Huỳnh Thành Duy", "Lục hồ sơ", "Đã lục giao lên scan", 120, ""],
    ]
    for r in data_nk:
        ws_nhat_ky.append(r)

    # --- 3. Tạo Sheet DASHBOARD ---
    ws_dashboard = wb.create_sheet(title="DASHBOARD")
    
    # Header DASHBOARD
    ws_dashboard.append(["BÁO CÁO TỔNG HỢP TIẾN ĐỘ CÔNG VIỆC"])
    ws_dashboard.append([])
    
    headers_db = ["Họ và Tên", "Nhiệm vụ", "Tổng khối lượng giao", "Khối lượng ĐÃ LÀM", "Khối lượng CÒN LẠI", "% Hoàn thành"]
    ws_dashboard.append(headers_db)
    
    # Dashboard Data (using formulas)
    for i, row in enumerate(data_mt):
        r_idx = i + 4 # starting from row 4
        name = row[1]
        task = row[3]
        total = row[4]
        
        # SUMIFS: =SUMIFS('NHAT KY'!E:E, 'NHAT KY'!B:B, name, 'NHAT KY'!C:C, task)
        # Note: openpyxl will write english formula names like SUMIFS
        sumifs_formula = f"=SUMIFS('NHAT KY'!E:E, 'NHAT KY'!B:B, A{r_idx}, 'NHAT KY'!C:C, B{r_idx})"
        remaining_formula = f"=C{r_idx}-D{r_idx}"
        pct_formula = f"=IF(C{r_idx}=0, 0, D{r_idx}/C{r_idx})"
        
        ws_dashboard.append([name, task, total, sumifs_formula, remaining_formula, pct_formula])
        # Format percentage
        ws_dashboard[f'F{r_idx}'].number_format = '0.00%'
        
    # --- Format Sheets ---
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    
    for ws in wb.worksheets:
        # Format header row
        header_row = 1 if ws.title != "DASHBOARD" else 3
        for cell in ws[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = min(adjusted_width, 40)
            
    # Save Workbook
    wb.save(file_path)
    print(f"Thành công: Đã tạo file {file_path}")

if __name__ == '__main__':
    create_excel_master('/Users/admin/Desktop/BINHDUONG/QUẢN LÝ CÔNG VIỆC KẾ HOẠCH 9999 - MASTER.xlsx')
